"""Regenerate the Word learning report with the bundled Python runtime."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "C++图书管理系统语法复习报告.docx"


def set_font(style, name="Microsoft YaHei", size=10.5, bold=False):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)


def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def border_table(table):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom", "left", "right"):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:w"), "100")
        tag.set(qn("w:type"), "dxa")
        mar.append(tag)
    tblpr.append(mar)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(2.1)
sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.2)
sec.right_margin = Cm(2.2)

normal = doc.styles["Normal"]
set_font(normal)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.34
set_font(doc.styles["Title"], size=20, bold=True)
doc.styles["Title"].paragraph_format.space_after = Pt(12)
title_ppr = doc.styles["Title"].element.get_or_add_pPr()
for old in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(old)
no_border = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "nil")
no_border.append(bottom)
title_ppr.append(no_border)
for name, size, before, after in (("Heading 1", 15, 14, 7), ("Heading 2", 12, 10, 5)):
    style = doc.styles[name]
    set_font(style, size=size, bold=True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True


def p(text="", style=None):
    return doc.add_paragraph(text, style)


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def bullet(text):
    para = doc.add_paragraph(style="List Bullet")
    para.add_run(text)
    para.paragraph_format.space_after = Pt(3)


def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Cm(0.45)
    para.paragraph_format.space_before = Pt(3)
    para.paragraph_format.space_after = Pt(6)
    for line in text.split("\n"):
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(40, 57, 72)
        run.add_break()


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Cm(width)
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Cm(widths[i])
        cell.text = header
        shade(cell, "233B53")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].width = Cm(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j % 2:
                shade(cells[i], "F4F7FA")
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.2
                for run in para.runs:
                    run.font.size = Pt(9)
    border_table(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def topic(title, location, usage, caution):
    para = doc.add_paragraph()
    para.paragraph_format.keep_together = True
    lead = para.add_run(title + "  ")
    lead.bold = True
    para.add_run("代码位置：" + location + "。")
    para.add_run("用法：" + usage + "。")
    para.add_run("注意：" + caution + "。")


doc.add_paragraph("C++17 图书管理系统语法复习报告", "Title")
p("以能运行的图书管理系统为主线，按建模、馆藏操作、文件持久化、命令行交互和验证的顺序复习 C++ 基础。阅读时建议先运行程序，再按本报告定位源码，最后独立完成练习。")
p("项目包含可执行主程序 library_manager.exe、测试程序 library_tests.exe、补充实验 syntax_lab.exe、完整源码和一键构建脚本。主程序管理图书条目及副本数量；借阅和归还只改变数量，不记录读者。")

h1("1 学习路线与运行方式")
table(["阶段", "实现目标", "主要语法与源码"], [
    ("1 图书建模", "定义合法图书与借还规则", "class、构造函数、const、异常；book.*"),
    ("2 馆藏管理", "新增、查找、删除、排序、统计", "vector、指针、lambda、算法；library.*"),
    ("3 文件持久化", "保存和恢复书库", "抽象类、虚函数、文件流、RAII；storage.*"),
    ("4 用户交互", "菜单和可靠的数字输入", "main、switch、while、模板；main.cpp"),
    ("5 验证", "跑通正常与边界场景", "assert、接口替身；test_library.cpp"),
], [2.2, 4.1, 9.3])
p("在 PowerShell 中进入 D:\\chatGPT\\list，运行以下命令。build.ps1 需要 g++ 支持 C++17；本机使用 MinGW-w64 g++ 8.1.0 构建通过。")
code(".\\build.ps1\n.\\library_tests.exe .\\test_catalog.db\n.\\library_manager.exe .\\books.db")
p("首次启动时可从空书库开始。程序每次成功修改后写入数据文件；重新运行并使用同一文件可恢复数据。命令行菜单使用英文，源码注释与本报告使用中文，便于避开不同 Windows 终端的中文编码差异。")

h1("2 业务模型与关键约束")
p("一条 Book 记录有编号、书名、作者、出版年、总册数和可借册数。编号大于零；书名和作者不能全为空白，也不能含换行；年份限定在 0 至 2100；总册数大于零；可借册数在 0 与总册数之间。借阅使可借册数减一，归还使其加一，越界操作返回 false。")
p("Library 拥有一个 CatalogSnapshot，其中 std::vector<Book> 存放记录，next_id 分配新编号。删除仅在该书所有副本已归还时成功。保存文件的首行为版本标识和下一个编号，后续每行一条记录；读取时检查格式、重复编号和对象约束。")
code('LIBRARY_V1 3\n1 "C++ Primer" "Lippman" 2012 2 1\n2 "Algorithms" "Sedgewick" 2011 1 1')

h1("3 第一阶段 图书类与基础表达式")
topic("头文件与作用域", "include/book.hpp、src/book.cpp", "#pragma once 避免同一翻译单元重复包含；namespace library 将类型放在独立命名空间；声明写头文件，定义写 .cpp", "每个 .cpp 都要包含自己用到的标准库头；不要在头文件写 using namespace std")
topic("类型与成员", "Book 的 int、std::string 和 private 成员", "int 存数量和编号，std::string 存可变长度文本；private 阻止外部绕过规则直接修改副本数", "int 有范围，文本编码与字节数不同；不能把显示字符数等同于 std::string::size()")
topic("构造函数与初始化列表", "Book::Book", "创建对象时在冒号后的初始化列表构造成员，并在函数体检查跨字段约束", "成员按声明顺序初始化，不按列表书写顺序；无效对象应抛异常而不是先创建后修补")
topic("值传递与移动", "Book::Book、Library::add_book", "std::string 按值接收，调用方可传左值或临时值；std::move 把参数内容转交给成员", "被移动的变量仍有效，但内容未指定；需要继续使用时不要依赖原来的文本")
topic("const 与引用", "Book::title、operator<<", "const 成员函数承诺不修改对象；const std::string& 返回只读引用，避免复制；输出函数以 const Book& 读取对象", "返回的引用只在 Book 存活且未被容器操作使之失效时有效；不要返回局部变量的引用")
topic("条件、比较与自增减", "Book::borrow、Book::return_copy", "if 检查边界，-- 和 ++ 更新计数，bool 表示业务操作是否成功", "先判断再修改，保证 0 ≤ available ≤ total；前缀和后缀在这里只会影响表达式值，单独使用优先选前缀")
topic("异常与 noexcept", "Book 构造函数、borrow、return_copy", "非法输入抛 std::invalid_argument；借还只做整数比较和修改，标注 noexcept", "noexcept 函数若抛出异常会终止程序；只在确实不会抛异常时使用")
topic("运算符重载", "operator<<(std::ostream&, const Book&)", "重载流插入运算符使 std::cout << book 可读，并返回流引用支持连续输出", "不要改变原有运算符的直觉语义；输出函数不应偷偷修改图书")

h1("4 第二阶段 容器 算法与查找")
topic("std::vector 与范围 for", "CatalogSnapshot::books、Library::statistics", "vector 自动管理连续存储；for (const auto& book : books) 遍历且不复制元素", "push_back、emplace_back 和 erase 可能使元素指针、引用、迭代器失效；不要长期缓存 find_by_id 返回值")
topic("auto 与类型推断", "Library::find_by_id、sorted_books", "auto 接收迭代器或容器副本，减少冗长类型名；const auto& 保留只读引用语义", "普通 auto 会复制并丢掉顶层 const/引用；需要避免复制时明确写 &")
topic("指针与 nullptr", "Library::find_by_id", "返回 Book* 表示可能找不到；找到时用 &*it 取得元素地址，未找到时返回 nullptr", "使用前必须判断指针；指针只借用容器元素，不拥有对象，也会随 vector 扩容或删除失效")
topic("函数重载与 const 正确性", "两个 find_by_id 重载", "非 const 对象能获得 Book*，const 对象只能获得 const Book*，编译器据对象常量性选重载", "不要用 const_cast 绕开限制；只读查询不能暴露可写指针")
topic("lambda 捕获", "find_by_id、remove_book、sorted_books", "[id] 与 [field] 把局部值复制进匿名函数，供 find_if、remove_if 和 sort 使用", "捕获引用时要确认回调执行期间被引用对象仍存活；比较函数要满足严格弱序")
topic("标准算法", "std::find_if、std::sort、std::remove_if", "查找、排序和删除判定交给标准库；删除使用 erase(remove_if(...), end()) 惯用法", "remove_if 只移动元素并返回新的逻辑末尾，必须再 erase 才真正缩短 vector")
topic("枚举类与 switch", "SortField、sorted_books", "enum class 把排序选项限制为 Id、Title、Year，switch 逐项选择比较规则", "强类型枚举不会自动转为 int；新增枚举值时更新所有处理分支")
topic("std::string_view", "Library::search", "查询参数只读且无需复制，可接收 std::string 或字符串字面量", "string_view 不拥有字符数据，不要保存一个指向临时字符串的视图；当前搜索是区分大小写的子串匹配")
topic("返回值与复制", "search、sorted_books", "返回 std::vector<Book> 的独立结果，调用者可安全遍历或排序，不依赖内部容器地址", "复制有成本；数据量很大时可设计只读视图，但必须明确生命周期规则")

h1("5 第三阶段 接口 多态与文件")
topic("struct 与 class", "CatalogSnapshot、Statistics、Book", "struct 默认公开，适合简单数据聚合；class 默认私有，适合维护不变量", "两者能力相同，区别只是默认访问级别和默认继承权限")
topic("抽象类与虚函数", "IStorage、TextFileStorage", "含纯虚函数的 IStorage 定义 load/save 协议；TextFileStorage 用 override 实现，可替换为测试内存存储", "基类析构函数必须为 virtual，才能经基类指针正确销毁派生对象；override 可帮助编译器发现签名写错")
topic("final 与显式构造", "TextFileStorage final、explicit 构造函数", "final 禁止继续继承该文件存储类；explicit 阻止字符串被隐式转换成存储对象", "final 只在确实不需要扩展该实现时使用；单参数构造函数通常考虑 explicit")
topic("智能指针与所有权", "Library::storage_、std::make_unique", "std::unique_ptr<IStorage> 表示 Library 独占存储对象；构造时 std::move 转移所有权", "移动后原 unique_ptr 为空；非拥有裸指针只可在对象存活期间观察，不能 delete")
topic("RAII 与文件流", "TextFileStorage::load/save", "std::ifstream 与 std::ofstream 离开作用域自动关闭文件；保存时仍主动 close 以检查写入错误", "打开失败、格式错误和写入失败要区分；当前覆盖写入不是原子操作，崩溃可能损坏文件")
topic("格式化输入输出", "std::quoted、std::istringstream", "std::quoted 可将含空格、引号的书名和作者写成可解析字段；逐行解析便于报告出错行", "数据格式有版本标识；从外部读取的数据仍必须经过构造校验，不能信任文件内容")
topic("异常传播", "storage.cpp、main.cpp", "底层抛 std::runtime_error，main 在最外层捕获 std::exception 并输出错误；普通业务失败用 bool", "异常适合无法继续的错误，不能用作常规分支；保存失败时程序退出，当前内存修改不会可靠落盘")

h1("6 第四阶段 命令行控制流与模板")
topic("main 参数与三目运算符", "main(int argc, char* argv[])", "argc 判断是否传入数据文件名，条件运算符选取命令行路径或默认 books.db", "argv[1] 只能在 argc > 1 时访问；默认相对路径相对于启动时工作目录")
topic("while、switch 与作用域", "main.cpp 菜单循环", "while (true) 持续处理用户操作；switch 按整数选项分派；case 内声明变量时用大括号建立局部作用域", "case 不会自动建立作用域，漏写 break 会落入下一分支；此程序用 return 结束菜单")
topic("函数模板", "read_number<T>", "同一输入逻辑处理菜单编号、图书编号、年份和册数，调用时 read_number<int>(...) 实例化", "模板定义通常需在使用处可见；此模板只用于具备流提取运算符的类型")
topic("输入流状态", "read_line、read_number", "统一用 getline 取得整行，再用 istringstream 解析一个数字并确认行尾只有空白", "混用 cin >> 数字与 getline 常残留换行；空输入或尾随字符必须拒绝")
topic("局部变量与初始化", "const int choice、T value{}", "用 const 限定不再改变的值，用 {} 将模板变量值初始化；减少未初始化读取风险", "花括号初始化还会拒绝一些窄化转换；变量作用域尽量缩小")

h1("7 验证结果与边界")
table(["检查场景", "期望结果", "验证位置"], [
    ("新增两本书", "编号依次为 1、2", "test_workflow"),
    ("借完全部副本再借", "第三次返回 false；统计借出 2 册", "test_workflow"),
    ("有未归还副本时删除", "删除被拒绝", "test_workflow"),
    ("归还后删除并重载", "删除成功，保存后重载仍不存在", "test_workflow"),
    ("空白书名", "构造函数抛 invalid_argument", "test_validation_and_file"),
    ("含引号书名的保存与恢复", "标题和可借册数保持一致", "test_validation_and_file"),
    ("文件中重复编号", "加载时抛 runtime_error", "test_validation_and_file"),
], [4.1, 6.5, 5.0])
p("在本机运行 build.ps1 后，主程序、测试程序和补充实验均使用 -std=c++17 -Wall -Wextra -Wpedantic 编译成功；library_tests.exe 输出 All tests passed。测试以断言检查业务规则，编译测试程序时不要定义 NDEBUG，否则 assert 会被禁用。")

h1("8 复习时应主动辨认的边界")
bullet("值、引用、指针：Book 值对象由 vector 拥有；const Book& 只借用并只读；Book* 用 nullptr 表示未找到。")
bullet("编译期与运行期：语法错误和类型不匹配由编译器检查；文件损坏、非法书目数据和写入失败需要运行期处理。")
bullet("资源生命周期：局部流对象与 unique_ptr 离开作用域自动释放；不要手动 delete 智能指针持有的对象。")
bullet("数据一致性：每次业务修改后保存，但覆盖写入遇到断电或异常退出仍可能不完整；生产系统应使用临时文件和可靠替换。")
bullet("功能范围：项目不记录借阅人、时间、逾期或并发修改；搜索区分大小写，删除不复用旧编号。")

h1("9 递进练习")
for number, title, task, point in [
    (1, "输入约束", "要求列表排序选项只能是 1、2、3，并在非法输入时重试。", "while、条件、函数复用"),
    (2, "分类字段", "给 Book 增加分类枚举，支持按分类筛选，并更新文件格式版本。", "enum class、构造函数、兼容性"),
    (3, "读者与借阅记录", "新增 Reader 和 Loan，将谁借了哪本书及日期保存下来。", "类之间的关系、vector、结构体"),
    (4, "异常安全保存", "先写临时文件，成功后再替换正式文件，并设计失败恢复步骤。", "RAII、文件错误、异常安全"),
    (5, "规模优化", "用 unordered_map 加速按编号查找，并分析与 vector 排序展示的取舍。", "哈希容器、迭代器、复杂度"),
]:
    para = doc.add_paragraph()
    para.add_run(f"练习 {number} {title}  ").bold = True
    para.add_run(task + " 重点：" + point + "。")

h1("10 补充语法实验")
p("运行 examples/syntax_lab.cpp 对应的 syntax_lab.exe，可查看三层书架的藏书量和可借比例。")
for label, usage, note in [
    ("constexpr 与 static_assert", "编译期固定书架数量并验证大于零", "运行期输入不能直接用于编译期断言"),
    ("std::array 与下标", "保存每层书架的册数，用 .at(i) 访问", ".at 检查范围，[] 不检查；下标从零开始"),
    ("std::map 与结构化绑定", "按编号查找书架，并用 [id, name] 遍历", "map 按键排序；find 返回 end 表示不存在"),
    ("std::optional", "find_shelf 返回字符串或 std::nullopt", "解引用之前先判断；返回的 string 拥有数据"),
    ("static_cast 与整数除法", "将 available 转为 double 再计算比例", "两个整数直接相除会截断小数"),
    ("using 与 accumulate", "ShelfId 是别名，accumulate 求和", "别名不是新类型；累加初值决定结果类型"),
]:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(3)
    para.paragraph_format.line_spacing = 1.18
    para.add_run(label + "  ").bold = True
    para.add_run("用法：" + usage + "。注意：" + note + "。")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("C++17 图书管理系统语法复习报告")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(90, 90, 90)

doc.save(OUT)
print(OUT)
